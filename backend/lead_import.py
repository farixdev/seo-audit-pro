# -*- coding: utf-8 -*-
"""
Lead-bank import pipeline for scraped lead files (typically Google Maps / GMB
exports where one file = one niche in one area).

Flow
  1. ``create_preview`` stores the uploaded file as a *pending* batch and
     returns headers, detected column mapping, sample rows, row count and a
     niche/area guess (from the file name and the data).
  2. ``start`` queues the batch; a background worker parses the file in
     chunks and, for every row:
       * no website              -> rejected (no_website)
       * unusable website        -> rejected (invalid_website)
       * repeated in this file   -> rejected (duplicate_in_file)
       * website or phone already in the bank -> rejected (duplicate_existing)
       * otherwise inserted into the lead bank
     Counters and rejected rows are committed with each chunk, so progress is
     exact and an interrupted import resumes where it stopped.
"""

import csv
import io
import json
import logging
import os
import re
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor

from backend import cache, config, db, events, realtime, security
from backend.errors import ApiError
from backend.util import normalize_company, normalize_domain, normalize_phone, now_iso

log = logging.getLogger("seo_audit_pro.import")

IMPORT_DIR = os.path.join(config.DATA_DIR, "imports")
os.makedirs(IMPORT_DIR, exist_ok=True)
MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_ROWS = 1_000_000
CHUNK = 1000
_EXECUTOR = ThreadPoolExecutor(max_workers=2, thread_name_prefix="lead-import")
_ACTIVE = set()
_ACTIVE_LOCK = threading.Lock()

# Lead fields a file column can map to.
FIELDS = ["company", "website", "phone", "email", "contact_name", "address", "category", "rating", "review_count",
          "maps_url", "city", "country"]

_CANDIDATES = {
    "website": ["website", "web site", "site", "domain", "homepage", "web", "website url", "company website"],
    "company": ["company", "company name", "business name", "business", "name", "title", "place name", "organization"],
    "phone": ["phone", "phone number", "phone_number", "telephone", "tel", "mobile", "contact number", "phones"],
    "email": ["email", "emails", "e mail", "mail", "contact email", "email address"],
    "contact_name": ["contact name", "contact person", "owner", "owner name", "full name"],
    "address": ["address", "full address", "street", "street address", "location address"],
    "category": ["category", "main category", "categories", "primary category", "type", "business type",
                 "subtitle", "industry"],
    "rating": ["rating", "total score", "stars", "average rating", "review rating", "score"],
    "review_count": ["reviews", "reviews count", "review count", "number of reviews", "user ratings total",
                     "reviews_count", "total reviews"],
    "maps_url": ["google maps url", "maps url", "place url", "google url", "url", "link", "maps link", "place link"],
    "city": ["city", "town", "locality", "area", "region", "state"],
    "country": ["country", "country code"],
}


def _norm(h):
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", str(h or ""))  # reviewsCount -> reviews Count
    return " ".join(re.sub(r"[_\-./]+", " ", text).strip().lower().split())


def detect_columns(headers):
    """Map lead fields to file columns (exact header match first, then word match)."""
    normed = [(h, _norm(h)) for h in headers if str(h or "").strip()]
    used, mapping = set(), {}
    order = ["website", "maps_url", "company", "phone", "email", "contact_name", "address", "category", "rating",
             "review_count", "city", "country"]
    for field in order:
        found = None
        for cand in _CANDIDATES[field]:
            found = next((h for h, n in normed if n == cand and h not in used), None)
            if found:
                break
        if not found and field not in ("maps_url",):
            for cand in _CANDIDATES[field]:
                found = next((h for h, n in normed if h not in used and cand in n.split()), None)
                if found:
                    break
        mapping[field] = found
        if found:
            used.add(found)
    # A generic "url" column is the website when no explicit website column exists.
    if not mapping["website"] and mapping["maps_url"] and "google" not in _norm(mapping["maps_url"]) \
            and "map" not in _norm(mapping["maps_url"]):
        mapping["website"], mapping["maps_url"] = mapping["maps_url"], None
    return mapping


_FILE_NOISE = re.compile(r"\b(gmb|google|maps?|scrape[ds]?|scraper|export|data|leads?|list|final|copy|\d{4}-\d{2}-\d{2}|"
                         r"\d{6,})\b", re.I)


_SMALL_WORDS = {"in", "of", "at", "on", "de", "la", "el", "le", "to", "by", "an", "or"}


def _title(text):
    """Title-case a guess; two-letter codes (TX, IL, UK, NY) stay upper-case."""
    text = re.sub(r"\s+", " ", text).strip(" -_,.")
    words = []
    for w in text.split():
        if len(w) == 2 and w.isalpha() and w.lower() not in _SMALL_WORDS:
            words.append(w.upper())
        elif w.isupper() and len(w) <= 3:
            words.append(w)
        else:
            words.append(w.capitalize())
    return " ".join(words)


def guess_niche_area(file_name, rows, mapping):
    """Best guess from names like "dentists_in_chicago_il.csv" or
    "Plumbers - London.xlsx", falling back to the most common category/city."""
    stem = os.path.splitext(os.path.basename(file_name or ""))[0]
    stem = re.sub(r"[_]+", " ", stem)
    stem = re.sub(r"\(\d+\)", " ", stem)
    niche = area = ""
    m = re.match(r"^(.*?)\s+(?:in|near|at)\s+(.+)$", stem, re.I)
    if m:
        niche, area = m.group(1), m.group(2)
    else:
        parts = [p for p in re.split(r"\s+[-|–]\s+|\s*,\s*", stem) if p.strip()]
        if len(parts) < 2 and "-" in stem:  # "roofers-dallas-tx"
            parts = [p for p in stem.split("-") if p.strip()]
        if len(parts) >= 2:
            niche, area = parts[0], " ".join(parts[1:])
    niche = _title(_FILE_NOISE.sub(" ", niche))
    area = _title(_FILE_NOISE.sub(" ", area))
    if not niche and mapping.get("category"):
        common = Counter((r.get(mapping["category"]) or "").strip() for r in rows if r.get(mapping["category"]))
        if common:
            niche = common.most_common(1)[0][0]
    if not area and mapping.get("city"):
        common = Counter((r.get(mapping["city"]) or "").strip() for r in rows if r.get(mapping["city"]))
        if common:
            area = common.most_common(1)[0][0]
    return niche[:80], area[:120]


# ---------------------------------------------------------------------------
# File reading (streaming)
# ---------------------------------------------------------------------------
def _decode_stream(path):
    with open(path, "rb") as fh:
        head = fh.read(65536)
    for enc in ("utf-8-sig", "cp1252"):
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def iter_rows(path):
    """Yield (headers, iterator of row dicts) for CSV/TSV/TXT/XLSX files."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        enc = _decode_stream(path)
        fh = open(path, encoding=enc, errors="replace", newline="")
        sample = fh.read(16384)
        fh.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = "\t" if "\t" in sample else (";" if sample.count(";") > sample.count(",") else ",")
        csv.field_size_limit(10 * 1024 * 1024)
        reader = csv.reader(fh, delimiter=delimiter)
        headers = [h.strip() for h in next(reader, [])]

        def gen():
            try:
                for values in reader:
                    if not any((v or "").strip() for v in values):
                        continue
                    yield {h: (values[i].strip() if i < len(values) else "") for i, h in enumerate(headers) if h}
            finally:
                fh.close()
        return headers, gen()
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        first = next(it, None) or []
        headers = [str(h).strip() if h is not None else "" for h in first]

        def gen():
            try:
                for values in it:
                    row = {h: ("" if values[i] is None else str(values[i]).strip())
                           for i, h in enumerate(headers) if h and i < len(values)}
                    if any(row.values()):
                        yield row
            finally:
                wb.close()
        return headers, gen()
    raise ApiError(422, "Choose a .csv, .tsv or .xlsx file. Older .xls files must be re-saved as .xlsx.")


# ---------------------------------------------------------------------------
# Preview (pending batch)
# ---------------------------------------------------------------------------
def create_preview(user, file_storage, niche="", area=""):
    if file_storage is None or not file_storage.filename:
        raise ApiError(422, "Choose a file to import.")
    name = os.path.basename(file_storage.filename.replace("\\", "/"))[:200]
    ext = os.path.splitext(name)[1].lower()
    if ext not in (".csv", ".tsv", ".txt", ".xlsx"):
        raise ApiError(422, f"{name}: only .csv, .tsv and .xlsx files can be imported.")
    batch_id = security.new_id("bat_")
    stored = f"{batch_id}{ext}"
    path = os.path.join(IMPORT_DIR, stored)
    size = 0
    with open(path, "wb") as out:
        while True:
            chunk = file_storage.stream.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_FILE_BYTES:
                out.close()
                os.remove(path)
                raise ApiError(413, f"{name} is larger than 200 MB. Split it into smaller files.")
            out.write(chunk)
    try:
        headers, rows = iter_rows(path)
        sample, total = [], 0
        for row in rows:
            total += 1
            if len(sample) < 200:
                sample.append(row)
            if total > MAX_ROWS:
                raise ApiError(422, f"{name} has more than {MAX_ROWS:,} rows. Split it into smaller files.")
    except ApiError:
        os.remove(path)
        raise
    except Exception:  # noqa: BLE001
        os.remove(path)
        raise ApiError(422, f"{name} could not be read. Check that it is a valid CSV or Excel file.")
    if not headers or total == 0:
        os.remove(path)
        raise ApiError(422, f"{name} has no data rows.")
    mapping = detect_columns(headers)
    guess_niche, guess_area = guess_niche_area(name, sample, mapping)
    niche = (niche or guess_niche).strip()[:80]
    area = (area or guess_area).strip()[:120]
    db.execute("""INSERT INTO lead_batches (id, file_name, stored_name, file_size, niche, area, niche_key, area_key, source,
                  status, headers, mapping, total_rows, created_by, created_at)
                  VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'pending', %s, %s, %s, %s, %s)""",
               (batch_id, name, stored, size, niche or None, area or None, (niche or "").lower() or None,
                (area or "").lower() or None, "File import", json.dumps(headers), json.dumps(mapping), total,
                user["id"], now_iso()))
    with_site = sum(1 for r in sample if mapping.get("website") and (r.get(mapping["website"]) or "").strip())
    return {"batch": batch_out(get_batch(batch_id)), "headers": headers, "mapping": mapping,
            "sample": sample[:6], "sample_with_website": with_site, "sample_size": len(sample)}


def get_batch(batch_id):
    return db.query_one("""SELECT b.*, u.name AS created_by_name, t.name AS team_name, a.name AS agent_name
                           FROM lead_batches b LEFT JOIN users u ON u.id = b.created_by
                           LEFT JOIN teams t ON t.id = b.team_id LEFT JOIN users a ON a.id = b.agent_id
                           WHERE b.id = %s""", (batch_id,))


def batch_out(row):
    if not row:
        return None
    keys = ("id", "file_name", "file_size", "niche", "area", "status", "total_rows", "processed_rows", "imported",
            "duplicates", "no_website", "invalid", "team_id", "agent_id", "team_name", "agent_name", "created_by",
            "created_by_name", "created_at", "started_at", "completed_at", "error")
    out = {k: row.get(k) for k in keys}
    out["rejected"] = int(row.get("duplicates") or 0) + int(row.get("no_website") or 0) + int(row.get("invalid") or 0)
    return out


# ---------------------------------------------------------------------------
# Processing
# ---------------------------------------------------------------------------
def _audience(batch):
    ids = {batch["created_by"], batch.get("agent_id"), events.team_lead_of(batch.get("team_id"))}
    ids.update(events.managers())
    return [i for i in ids if i]


def start(batch_id):
    db.execute("UPDATE lead_batches SET status = 'queued' WHERE id = %s AND status = 'pending'", (batch_id,))
    _submit(batch_id)


def _submit(batch_id):
    with _ACTIVE_LOCK:
        if batch_id in _ACTIVE:
            return
        _ACTIVE.add(batch_id)
    _EXECUTOR.submit(_run_safe, batch_id)


def resume_interrupted():
    """Continue imports that were running when the server stopped."""
    for row in db.query("SELECT id FROM lead_batches WHERE status IN ('queued', 'processing') ORDER BY created_at"):
        _submit(row["id"])


def _run_safe(batch_id):
    try:
        _run(batch_id)
    except Exception:  # noqa: BLE001
        log.exception("Import %s failed", batch_id)
        db.execute("UPDATE lead_batches SET status = 'failed', error = %s, completed_at = %s WHERE id = %s",
                   ("The import stopped because of an unexpected error. Already imported rows were kept.", now_iso(),
                    batch_id))
        batch = get_batch(batch_id)
        if batch:
            realtime.publish(_audience(batch), {"type": "import_progress", "batch": batch_out(batch)})
    finally:
        with _ACTIVE_LOCK:
            _ACTIVE.discard(batch_id)


def _text(value, limit):
    v = (value or "").strip()
    return v[:limit] if v else None


def _number(value, cast):
    v = (value or "").strip()
    if cast is float and "," in v and "." not in v and re.fullmatch(r"\d+,\d{1,2}", v):
        v = v.replace(",", ".")  # European decimal comma: "4,9"
    v = re.sub(r"[^\d.]", "", v.replace(",", ""))
    if not v:
        return None
    try:
        return cast(float(v))
    except ValueError:
        return None


def _rating(value):
    r = _number(value, float)
    return r if r is not None and 0 <= r <= 5 else None


def _prepare(raw, mapping, batch, now):
    """Turn a file row into a lead tuple, or (None, reason)."""
    get = lambda f: raw.get(mapping[f]) if mapping.get(f) else None  # noqa: E731
    website = (get("website") or "").strip()
    if not website:
        return None, "no_website"
    domain = normalize_domain(website.split()[0])
    if not domain or domain.endswith(("google.com", "facebook.com", "instagram.com", "goo.gl", "g.page")):
        return None, "invalid_website"
    company = _text(get("company"), 200) or domain
    email = (get("email") or "").split(",")[0].strip().lower()[:254] or None
    if email and "@" not in email:
        email = None
    phone = _text(get("phone"), 40)
    known = {c for c in mapping.values() if c}
    extra = {k: v for k, v in raw.items() if k not in known and v}
    area = batch["area"] or _text(get("city"), 120)
    lead = {
        "id": security.new_id("lead_"), "company": company,
        "website": website if website.lower().startswith(("http://", "https://")) else f"https://{domain}",
        "normalized_domain": domain, "contact_name": _text(get("contact_name"), 120), "email": email,
        "email_key": email, "phone": phone, "phone_key": normalize_phone(phone) or None,
        "company_key": normalize_company(company), "niche": batch["niche"] or _text(get("category"), 80),
        "country": _text(get("country"), 80), "area": area, "source": f"Import: {batch['file_name']}"[:80],
        "status": "assigned" if batch.get("agent_id") else "new", "assigned_team_id": batch.get("team_id"),
        "assigned_agent_id": batch.get("agent_id"), "created_by": batch["created_by"], "created_at": now,
        "updated_at": now, "address": _text(get("address"), 300), "category": _text(get("category"), 120),
        "rating": _rating(get("rating")), "review_count": _number(get("review_count"), int),
        "maps_url": _text(get("maps_url"), 500), "extra": json.dumps(extra, ensure_ascii=False)[:4000] if extra else None,
        "batch_id": batch["id"],
    }
    lead["niche_key"] = (lead["niche"] or "").lower() or None
    lead["area_key"] = (lead["area"] or "").lower() or None
    return lead, None


_LEAD_COLUMNS = ["id", "company", "website", "normalized_domain", "contact_name", "email", "email_key", "phone",
                 "phone_key", "company_key", "niche", "niche_key", "country", "area", "area_key", "source", "status",
                 "assigned_team_id", "assigned_agent_id", "created_by", "created_at", "updated_at", "address",
                 "category", "rating", "review_count", "maps_url", "extra", "batch_id"]
_REJECT_COLUMNS = ["id", "batch_id", "row_number", "company", "website", "phone", "reason", "existing_lead_id"]


def _commit_chunk(batch, rows, seen_domains, seen_phones, first_row_number):
    """Classify and store one chunk atomically. Returns counter deltas."""
    now = now_iso()
    for attempt in range(3):
        try:
            with db.transaction() as tx:
                prepared, rejects = [], []
                counts = {"imported": 0, "duplicates": 0, "no_website": 0, "invalid": 0}
                candidates = []
                local_domains, local_phones = set(), set()
                for offset, raw in enumerate(rows):
                    number = first_row_number + offset
                    lead, reason = _prepare(raw, batch["mapping"], batch, now)
                    if reason:
                        rejects.append((number, raw, reason, None))
                        continue
                    if lead["normalized_domain"] in seen_domains or lead["normalized_domain"] in local_domains or \
                            (lead["phone_key"] and (lead["phone_key"] in seen_phones or lead["phone_key"] in local_phones)):
                        rejects.append((number, raw, "duplicate_in_file", None))
                        continue
                    local_domains.add(lead["normalized_domain"])
                    if lead["phone_key"]:
                        local_phones.add(lead["phone_key"])
                    candidates.append((number, raw, lead))
                existing_domain, existing_phone = {}, {}
                if candidates:
                    domains = [c[2]["normalized_domain"] for c in candidates]
                    ph = ",".join(["%s"] * len(domains))
                    for r in tx.query(f"SELECT id, normalized_domain FROM leads WHERE normalized_domain IN ({ph})", domains):
                        existing_domain[r["normalized_domain"]] = r["id"]
                    phones = [c[2]["phone_key"] for c in candidates if c[2]["phone_key"]]
                    if phones:
                        ph = ",".join(["%s"] * len(phones))
                        for r in tx.query(f"SELECT id, phone_key FROM leads WHERE phone_key IN ({ph})", phones):
                            existing_phone.setdefault(r["phone_key"], r["id"])
                for number, raw, lead in candidates:
                    dup = existing_domain.get(lead["normalized_domain"]) or \
                        (existing_phone.get(lead["phone_key"]) if lead["phone_key"] else None)
                    if dup:
                        rejects.append((number, raw, "duplicate_existing", dup))
                    else:
                        prepared.append((number, raw, lead))
                inserted = tx.bulk_insert("leads", _LEAD_COLUMNS,
                                          [[lead[c] for c in _LEAD_COLUMNS] for _n, _r, lead in prepared],
                                          conflict_column="normalized_domain")
                for number, raw, lead in prepared:
                    if lead["id"] not in inserted:  # inserted concurrently by another import
                        rejects.append((number, raw, "duplicate_existing", None))
                counts["imported"] = len(inserted)
                reject_rows = []
                m = batch["mapping"]
                for number, raw, reason, dup in rejects:
                    key = {"no_website": "no_website", "invalid_website": "invalid"}.get(reason, "duplicates")
                    counts[key] += 1
                    reject_rows.append([security.new_id("rej_"), batch["id"], number,
                                        _text(raw.get(m["company"]) if m.get("company") else None, 200),
                                        _text(raw.get(m["website"]) if m.get("website") else None, 300),
                                        _text(raw.get(m["phone"]) if m.get("phone") else None, 40), reason, dup])
                tx.bulk_insert("lead_batch_rejects", _REJECT_COLUMNS, reject_rows)
                tx.execute("""UPDATE lead_batches SET processed_rows = processed_rows + %s, imported = imported + %s,
                                  duplicates = duplicates + %s, no_website = no_website + %s, invalid = invalid + %s
                              WHERE id = %s""",
                           (len(rows), counts["imported"], counts["duplicates"], counts["no_website"], counts["invalid"],
                            batch["id"]))
            seen_domains.update(local_domains)
            seen_phones.update(local_phones)
            return counts
        except db.IntegrityError:
            if attempt == 2:
                raise
            continue  # a concurrent import inserted the same website; reclassify


def _run(batch_id):
    batch = get_batch(batch_id)
    if not batch or batch["status"] not in ("queued", "processing"):
        return
    batch["mapping"] = json.loads(batch["mapping"] or "{}")
    path = os.path.join(IMPORT_DIR, batch["stored_name"] or "")
    if not os.path.isfile(path):
        db.execute("UPDATE lead_batches SET status = 'failed', error = %s, completed_at = %s WHERE id = %s",
                   ("The uploaded file is no longer available. Upload it again.", now_iso(), batch_id))
        return
    db.execute("UPDATE lead_batches SET status = 'processing', started_at = COALESCE(started_at, %s) WHERE id = %s",
               (now_iso(), batch_id))
    audience = _audience(batch)
    skip = int(batch["processed_rows"] or 0)
    seen_domains, seen_phones = set(), set()
    _headers, rows = iter_rows(path)
    chunk, number = [], 0
    first = None
    for row in rows:
        number += 1
        if number <= skip:
            # Resuming: rebuild the in-file duplicate sets from rows already processed.
            lead, reason = _prepare(row, batch["mapping"], batch, "")
            if not reason:
                seen_domains.add(lead["normalized_domain"])
                if lead["phone_key"]:
                    seen_phones.add(lead["phone_key"])
            continue
        if first is None:
            first = number
        chunk.append(row)
        if len(chunk) >= CHUNK:
            _commit_chunk(batch, chunk, seen_domains, seen_phones, first)
            cache.bump()
            chunk, first = [], None
            realtime.publish(audience, {"type": "import_progress", "batch": batch_out(get_batch(batch_id))})
    if chunk:
        _commit_chunk(batch, chunk, seen_domains, seen_phones, first)
    cache.bump()  # before "completed" is visible, so nobody reads stale counts after seeing it
    db.execute("UPDATE lead_batches SET status = 'completed', completed_at = %s WHERE id = %s", (now_iso(), batch_id))
    cache.bump()
    done = get_batch(batch_id)
    creator = db.query_one("SELECT * FROM users WHERE id = %s", (batch["created_by"],))
    events.audit(creator, "LEADS_IMPORTED", "lead_batch", batch_id, {
        "file": batch["file_name"], "niche": batch["niche"], "area": batch["area"], "rows": done["processed_rows"],
        "imported": done["imported"], "duplicates": done["duplicates"], "no_website": done["no_website"],
        "invalid": done["invalid"], "team_id": batch.get("team_id"), "agent_id": batch.get("agent_id")})
    realtime.publish(audience, {"type": "import_progress", "batch": batch_out(done)})
    summary = (f"{done['imported']:,} imported · {done['duplicates']:,} duplicates · "
               f"{done['no_website']:,} without website removed")
    events.notify([batch["created_by"]], "import_completed", f"Import finished: {batch['file_name']}", summary,
                  "lead_batch", batch_id)
    if batch.get("agent_id") and done["imported"] and batch["agent_id"] != batch["created_by"]:
        events.notify([batch["agent_id"]], "lead_assigned", "Leads assigned",
                      f"{done['imported']:,} {batch['niche'] or ''} leads in {batch['area'] or 'a new area'} were "
                      "assigned to you.", "lead_batch", batch_id)
    try:
        os.remove(path)  # the leads and rejected rows are stored; the raw file is no longer needed
        db.execute("UPDATE lead_batches SET stored_name = NULL WHERE id = %s", (batch_id,))
    except OSError:
        pass


def rejects_csv(batch_id):
    """Rejected rows of a batch as CSV text."""
    rows = db.query("""SELECT r.row_number, r.company, r.website, r.phone, r.reason, l.company AS existing_company,
                              l.normalized_domain AS existing_domain
                       FROM lead_batch_rejects r LEFT JOIN leads l ON l.id = r.existing_lead_id
                       WHERE r.batch_id = %s ORDER BY r.row_number""", (batch_id,))
    labels = {"no_website": "No website", "invalid_website": "Website not usable",
              "duplicate_in_file": "Duplicate within this file", "duplicate_existing": "Already in lead bank"}
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Row", "Company", "Website", "Phone", "Reason", "Existing lead", "Existing website"])
    for r in rows:
        w.writerow([r["row_number"], r["company"] or "", r["website"] or "", r["phone"] or "",
                    labels.get(r["reason"], r["reason"]), r["existing_company"] or "", r["existing_domain"] or ""])
    return buf.getvalue()


def cleanup_pending(max_age_seconds=86400):
    """Delete previews that were never imported."""
    from backend.util import iso_ago
    for row in db.query("SELECT id, stored_name FROM lead_batches WHERE status = 'pending' AND created_at < %s",
                        (iso_ago(max_age_seconds),)):
        if row["stored_name"]:
            try:
                os.remove(os.path.join(IMPORT_DIR, row["stored_name"]))
            except OSError:
                pass
        db.execute("DELETE FROM lead_batches WHERE id = %s AND status = 'pending'", (row["id"],))
